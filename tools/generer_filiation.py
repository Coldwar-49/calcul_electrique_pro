#!/usr/bin/env python3
"""Génère lib/core/donnees/filiation_merlin_gerin.dart depuis filiation.xlsx.

Usage : python tools/generer_filiation.py docs/filiation.xlsx

Chaque feuille numérotée du classeur contient un ou plusieurs blocs :
  - une ligne « Amont » suivie des appareils amont (en colonnes) ;
  - une ligne « Pdc » avec le pouvoir de coupure de chaque appareil amont ;
  - une ligne « Aval » puis un appareil aval par ligne, dont les cellules
    donnent le pouvoir de coupure renforcé (kA) de l'aval protégé par l'amont ;
    une cellule vide = pas de filiation.
Le script ne devine rien : une cellule non numérique, un bloc incomplet ou une
feuille sans bloc est signalé sur la sortie d'erreur.
"""
import re
import sys

import openpyxl

TITRE = re.compile(r'^Réseau\s+(\d+)\s*V\s*-\s*(.*?)\s*-\s*Année\s+(.*)$')


def texte(v):
    """Texte d'une cellule ; plusieurs lignes (ex. « NS2000N / NS2500N ») sont
    jointes par « / » pour rester sur une seule ligne dans le fichier Dart."""
    if v is None:
        return None
    lignes = [l.strip() for l in str(v).splitlines() if l.strip()]
    return ' / '.join(lignes) if lignes else None


def trouver_titre(ws, ligne, col):
    """Titre du bloc : texte « Réseau … » dans la colonne du bloc, au-dessus."""
    for r in range(ligne - 1, max(ligne - 5, 0), -1):
        t = texte(ws.cell(r, col).value)
        if t and t.startswith('Réseau'):
            return t
    return None


def lire_bloc(ws, ligne, col, avertir):
    # Appareils amont : cellules contiguës à droite de « Amont ».
    amonts, c = [], col + 1
    while texte(ws.cell(ligne, c).value):
        amonts.append(texte(ws.cell(ligne, c).value))
        c += 1
    if not amonts:
        avertir(f'bloc {ws.cell(ligne, col).coordinate} sans appareil amont')
        return None

    # En-têtes sur plusieurs lignes : les cellules situées sous un nom d'amont,
    # avant la ligne « Pdc » ou « Aval », complètent ce nom (ex. « Masterpact »
    # puis « NT L1 » donnent « Masterpact / NT L1 »).
    fin_entete = None
    for r in range(ligne + 1, ligne + 8):
        t = texte(ws.cell(r, col).value)
        if t and t.lower() in ('pdc', 'aval'):
            fin_entete = r
            break
    if fin_entete is None:
        avertir(f'bloc {ws.cell(ligne, col).coordinate} sans ligne « Pdc » ni « Aval »')
        return None
    for i in range(len(amonts)):
        for r in range(ligne + 1, fin_entete):
            complement = texte(ws.cell(r, col + 1 + i).value)
            if complement:
                amonts[i] = f'{amonts[i]} / {complement}'

    # Ligne « Pdc » (facultative) et ligne « Aval ».
    pdc, ligne_aval = [None] * len(amonts), None
    for r in range(ligne + 1, ligne + 8):
        t = texte(ws.cell(r, col).value)
        if t and t.lower() == 'pdc':
            for i in range(len(amonts)):
                v = ws.cell(r, col + 1 + i).value
                if isinstance(v, (int, float)):
                    pdc[i] = v
                elif v is not None:
                    avertir(f'{ws.cell(r, col + 1 + i).coordinate}: Pdc non numérique {v!r}')
        if t and t.lower() == 'aval':
            ligne_aval = r
            break
    if ligne_aval is None:
        avertir(f'bloc {ws.cell(ligne, col).coordinate} sans ligne « Aval »')
        return None

    avals, valeurs = [], []
    r = ligne_aval + 1
    while texte(ws.cell(r, col).value):
        nom = texte(ws.cell(r, col).value)
        ligne_vals = []
        for i in range(len(amonts)):
            cell = ws.cell(r, col + 1 + i)
            v = cell.value
            if v is None or (isinstance(v, str) and not v.strip()):
                ligne_vals.append(None)
            elif isinstance(v, (int, float)):
                ligne_vals.append(v)
            else:
                avertir(f'{cell.coordinate}: valeur non numérique {v!r} (ignorée)')
                ligne_vals.append(None)
        avals.append(nom)
        valeurs.append(ligne_vals)
        r += 1
    if not avals:
        avertir(f'bloc {ws.cell(ligne, col).coordinate} sans appareil aval')
        return None
    return amonts, pdc, avals, valeurs


def dart_str(s):
    return "'" + s.replace('\\', '\\\\').replace("'", "\\'").replace('$', '\\$') + "'"


def dart_num(v):
    if v is None:
        return 'null'
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v)


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    tables, ignorees = [], []

    for ws in wb.worksheets:
        def avertir(msg, nom=ws.title):
            print(f'  [feuille {nom}] {msg}', file=sys.stderr)

        ancres = [(c.row, c.column) for row in ws.iter_rows() for c in row
                  if texte(c.value) and texte(c.value).lower() == 'amont']
        if not ancres:
            ignorees.append(ws.title)
            continue
        for ligne, col in sorted(ancres, key=lambda a: (a[1], a[0])):
            titre = trouver_titre(ws, ligne, col)
            m = TITRE.match(titre) if titre else None
            if not m:
                avertir(f'titre introuvable ou illisible pour {ws.cell(ligne, col).coordinate}: {titre!r}')
                continue
            bloc = lire_bloc(ws, ligne, col, avertir)
            if bloc is None:
                continue
            amonts, pdc, avals, valeurs = bloc
            tables.append(dict(
                numero=int(ws.title), annee=m.group(3).strip(),
                tension=int(m.group(1)), libelle=m.group(2).strip(),
                amonts=amonts, pdc=pdc, avals=avals, valeurs=valeurs))

    # Rapport.
    print(f'{len(tables)} tables lues dans {len(wb.worksheets)} feuilles.', file=sys.stderr)
    if ignorees:
        print(f'Feuilles sans bloc « Amont » (non portées) : {", ".join(ignorees)}', file=sys.stderr)
    nb_val = sum(1 for t in tables for l in t['valeurs'] for v in l if v is not None)
    print(f'{nb_val} valeurs de pouvoir de coupure renforcé.', file=sys.stderr)

    # Fichier Dart.
    out = ["/// Tableaux de filiation Merlin Gerin / Schneider (1998 à 2012).",
           "///",
           "/// FICHIER GÉNÉRÉ par tools/generer_filiation.py depuis filiation.xlsx : ne pas",
           "/// modifier à la main. Valeurs en kA ; `null` = pas de filiation.",
           "library;",
           "",
           "import 'filiation.dart';",
           "",
           "const List<TableFiliation> tablesFiliationMerlinGerin = ["]
    for t in tables:
        out.append('  TableFiliation(')
        out.append(f"    numero: {t['numero']},")
        out.append(f"    catalogue: {dart_str(t['annee'])},")
        out.append(f"    tension: {t['tension']},")
        out.append(f"    libelle: {dart_str(t['libelle'])},")
        out.append('    amonts: [' + ', '.join(dart_str(a) for a in t['amonts']) + '],')
        out.append('    pdcAmont: <double?>[' + ', '.join(dart_num(v) for v in t['pdc']) + '],')
        out.append('    avals: [' + ', '.join(dart_str(a) for a in t['avals']) + '],')
        out.append('    valeurs: <List<double?>>[')
        for ligne in t['valeurs']:
            out.append('      <double?>[' + ', '.join(dart_num(v) for v in ligne) + '],')
        out.append('    ],')
        out.append('  ),')
    out.append('];')
    out.append('')
    chemin = 'lib/core/donnees/filiation_merlin_gerin.dart'
    with open(chemin, 'w', encoding='utf-8', newline='\n') as f:
        f.write('\n'.join(out))
    print(f'Écrit : {chemin}', file=sys.stderr)


if __name__ == '__main__':
    main()
